#!/usr/bin/env python3
"""Regression: Claude Usage-style decks must not land on wrong wells or overlay chrome."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
sys.path.insert(0, str(ROOT))

from colorways import is_layout_chrome  # noqa: E402
from helpers import blob_hash, is_decorative_icon  # noqa: E402
from test_margins_images import TEMPLATE, _convert, _raster_hashes, _template_content_photo_hashes  # noqa: E402
from test_overlap_layers import LEFTOVER_DUMMY, _assert_no_dummy, _iter_text  # noqa: E402

EMU = 914400


def _icon_png(color: tuple[int, int, int] = (16, 80, 64)) -> bytes:
    im = Image.new("RGBA", (256, 256), (0, 0, 0, 0))
    draw = ImageDraw.Draw(im)
    draw.rounded_rectangle((40, 32, 216, 224), radius=18, fill=color + (255,))
    draw.rectangle((88, 88, 168, 104), fill=(255, 255, 255, 255))
    draw.rectangle((88, 120, 168, 136), fill=(255, 255, 255, 255))
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _write_claude_like_source(path: Path) -> None:
    """PptxGenJS-style 16:9 deck matching the uploaded green convert's content."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    icon = _icon_png()

    # Cover: org chrome above three side-by-side title boxes + empty decorative cards.
    s = prs.slides.add_slide(blank)
    org = s.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(6.0), Inches(0.4))
    org.text_frame.paragraphs[0].text = "EVERGLADES FOUNDATION"
    for i, text in enumerate(("Claude Usage", "Executive Summary", "60-Day Adoption Review  •  Prepared for Leadership")):
        box = s.shapes.add_textbox(Inches(0.5 + i * 4.2), Inches(2.8), Inches(3.9), Inches(1.4))
        box.text_frame.paragraphs[0].text = text
        deco = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5 + i * 4.2), Inches(2.4), Inches(3.9), Inches(3.2))
        deco.fill.solid()
        deco.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xF7, 0xF9, 0xEA)

    # Overview KPIs with four 256px icons.
    s = prs.slides.add_slide(blank)
    title = s.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.2), Inches(0.6))
    title.text_frame.paragraphs[0].text = "Overview"
    intro = s.shapes.add_textbox(Inches(0.5), Inches(1.05), Inches(12.2), Inches(0.7))
    intro.text_frame.paragraphs[0].text = (
        "Weekly active usage grew 30% week-over-week, with 13 of 19 seats active. "
        "Usage is concentrated in document, research, and presentation workflows."
    )
    kpis = [
        ("13 / 19", "Weekly Active Seats"),
        ("WoW Growth", "Cowork Sessions"),
        ("Conversations", "File Operations"),
        ("~20 hrs", "Time Saved"),
    ]
    for i, (head, sub) in enumerate(kpis):
        x = 0.5 + i * 3.2
        s.shapes.add_picture(io.BytesIO(icon), Inches(x + 0.9), Inches(2.0), Inches(1.1), Inches(1.1))
        h = s.shapes.add_textbox(Inches(x), Inches(3.3), Inches(3.0), Inches(0.4))
        h.text_frame.paragraphs[0].text = head
        b = s.shapes.add_textbox(Inches(x), Inches(3.75), Inches(3.0), Inches(0.4))
        b.text_frame.paragraphs[0].text = sub
        card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.9), Inches(3.0), Inches(2.5))
        card.fill.solid()
        card.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xEE, 0xF2, 0xD8)

    # Who's using: three columns + doughnut + tables.
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.2), Inches(0.5))
    t.text_frame.paragraphs[0].text = "Who's Using Claude"
    cols = [
        ("Usage is led by the Science team, with a small group of power users driving most activity.", "BY GROUP"),
        ("SEAT UTILIZATION", "13 of 19"),
        ("TOP USERS BY CHATS", ""),
    ]
    for i, (a, b) in enumerate(cols):
        x = 0.5 + i * 4.2
        ta = s.shapes.add_textbox(Inches(x), Inches(1.0), Inches(3.9), Inches(0.8))
        ta.text_frame.paragraphs[0].text = a
        if b:
            tb = s.shapes.add_textbox(Inches(x), Inches(1.85), Inches(3.9), Inches(0.4))
            tb.text_frame.paragraphs[0].text = b
    data = CategoryChartData()
    data.categories = ["Active", "Inactive"]
    data.add_series("Seats", (13, 6))
    s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.5), Inches(2.5), Inches(4.2), Inches(3.4), data)
    table = s.shapes.add_table(3, 3, Inches(5.0), Inches(2.5), Inches(7.8), Inches(1.6)).table
    for r, row in enumerate((("Group", "Active", "Messages"), ("Science", "7", "1,280"), ("Admin", "1", "640"))):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    table2 = s.shapes.add_table(3, 2, Inches(5.0), Inches(4.3), Inches(7.8), Inches(1.6)).table
    for r, row in enumerate((("Member", "Chats"), ("yacosta@evergladesfoundation.org", "229"), ("lkumar@evergladesfoundation.org", "97"))):
        for c, val in enumerate(row):
            table2.cell(r, c).text = val

    # Two points + bar chart + table + icon.
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.2), Inches(0.5))
    t.text_frame.paragraphs[0].text = "How They're Using Claude"
    left = s.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(6.0), Inches(1.2))
    left.text_frame.paragraphs[0].text = "Skill usage skews toward document creation and research; connector adoption is still early-stage."
    left2 = s.shapes.add_textbox(Inches(0.5), Inches(2.3), Inches(6.0), Inches(0.4))
    left2.text_frame.paragraphs[0].text = "SKILL USES"
    right = s.shapes.add_textbox(Inches(7.0), Inches(1.1), Inches(5.8), Inches(0.4))
    right.text_frame.paragraphs[0].text = "CONNECTORS"
    right2 = s.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.8), Inches(1.2))
    right2.text_frame.paragraphs[0].text = (
        "Connector adoption remains concentrated among a few users. Expanding Microsoft 365, Asana, and Chrome usage across more team members is a clear growth opportunity."
    )
    bars = CategoryChartData()
    bars.categories = ["docx", "read-paper", "pptx", "pdf", "co-author"]
    bars.add_series("Uses", (54, 36, 29, 24, 22))
    s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(3.0), Inches(7.5), Inches(3.8), bars)
    tbl = s.shapes.add_table(3, 2, Inches(8.3), Inches(3.2), Inches(4.5), Inches(1.5)).table
    for r, row in enumerate((("Connector", "Users"), ("Visualize", "5"), ("Microsoft 365", "1"))):
        for c, val in enumerate(row):
            tbl.cell(r, c).text = val
    s.shapes.add_picture(io.BytesIO(icon), Inches(11.6), Inches(6.2), Inches(1.0), Inches(1.0))

    # Agentic KPIs with four icons.
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.2), Inches(0.5))
    t.text_frame.paragraphs[0].text = "How Agentic Is the Work"
    items = [
        ("14.2", "Actions per Prompt"),
        ("Org-wide average", "Claude Code | Tool calls per message"),
        ("Cowork", "Actions per prompt | Chat"),
        ("Tool uses per message", "Note: 0 Claude Code PRs created to date."),
    ]
    for i, (a, b) in enumerate(items):
        x = 0.5 + i * 3.2
        s.shapes.add_picture(io.BytesIO(icon), Inches(x + 0.9), Inches(1.4), Inches(1.1), Inches(1.1))
        ha = s.shapes.add_textbox(Inches(x), Inches(2.7), Inches(3.0), Inches(0.4))
        ha.text_frame.paragraphs[0].text = a
        hb = s.shapes.add_textbox(Inches(x), Inches(3.15), Inches(3.0), Inches(0.8))
        hb.text_frame.paragraphs[0].text = b

    # Results + table + lightning icon.
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(7.5), Inches(0.5))
    t.text_frame.paragraphs[0].text = "Results & Time Saved"
    body = s.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(7.5), Inches(1.6))
    body.text_frame.paragraphs[0].text = (
        "Verified output counts × minutes-per-task assumptions. ~20 HOURS SAVED. "
        "Conversations account for the majority of time saved (93%)."
    )
    tbl = s.shapes.add_table(4, 4, Inches(0.5), Inches(3.0), Inches(7.5), Inches(2.4)).table
    rows = (
        ("Verified Output", "Volume", "Min / Each", "Total Minutes"),
        ("File operations", "55", "1", "55"),
        ("Conversations", "283", "4", "1,132"),
        ("Designs created", "1", "30", "30"),
    )
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            tbl.cell(r, c).text = val
    s.shapes.add_picture(io.BytesIO(icon), Inches(9.2), Inches(1.5), Inches(2.8), Inches(2.8))

    # Three takeaways.
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.2), Inches(0.5))
    t.text_frame.paragraphs[0].text = "Key Takeaways"
    takes = [
        (
            "Adoption is growing quickly (+30% WoW), but usage is concentrated: 68% seat utilization, with one user driving the majority of volume.",
            "The Science team is the primary adopter — broader rollout to other groups is an untapped opportunity.",
        ),
        (
            "Document and research skills (docx, read-paper, pptx, pdf) are the clear early use cases.",
            "Connector and agentic workflows (Asana, Microsoft 365, Claude Code) remain early-stage relative to document work.",
        ),
        (
            "Even at current adoption, Claude has returned ~20 hours of time savings in 60 days — with clear room to scale.",
            "",
        ),
    ]
    for i, (a, b) in enumerate(takes):
        x = 0.5 + i * 4.2
        ta = s.shapes.add_textbox(Inches(x), Inches(1.2), Inches(3.9), Inches(1.6))
        ta.text_frame.paragraphs[0].text = a
        if b:
            tb = s.shapes.add_textbox(Inches(x), Inches(2.9), Inches(3.9), Inches(1.4))
            tb.text_frame.paragraphs[0].text = b
        hair = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(5.8), Inches(3.9), Emu(0))
        hair.fill.solid()
        hair.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xC1, 0xD4, 0x51)

    prs.save(str(path))


def _picture_shapes(slide):
    return [s for s in slide.shapes if s.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}]


def _overlap(a, b) -> float:
    aw, ah = max(1, int(a.width or 0)), max(1, int(a.height or 0))
    left = max(int(a.left or 0), int(b.left or 0))
    top = max(int(a.top or 0), int(b.top or 0))
    right = min(int(a.left or 0) + aw, int(b.left or 0) + int(b.width or 0))
    bottom = min(int(a.top or 0) + ah, int(b.top or 0) + int(b.height or 0))
    if right <= left or bottom <= top:
        return 0.0
    return ((right - left) * (bottom - top)) / (aw * ah)


class IconDetectTests(unittest.TestCase):
    def test_256_png_is_icon_large_photo_is_not(self):
        self.assertTrue(is_decorative_icon({"image": {"px": [256, 256], "bytes": 4000, "ext": "png"}}))
        self.assertFalse(is_decorative_icon({"image": {"px": [400, 600], "bytes": 80000, "ext": "png"}}))
        self.assertFalse(is_decorative_icon({"image": {"px": [1600, 900], "bytes": 220000, "ext": "jpg"}}))
        self.assertFalse(
            is_decorative_icon({"image": {"px": [120, 120], "bytes": 237, "ext": "png"}, "area_ratio": 0.46}),
            "a low-res blob that already fills a photo well is still a photograph",
        )
        self.assertTrue(
            is_decorative_icon({"image": {"px": [256, 256], "bytes": 4000, "ext": "png"}, "area_ratio": 0.10}),
            "a 256px UI glyph scaled onto a KPI card is still an icon, not a photo well",
        )


@unittest.skipUnless(TEMPLATE.exists(), "official template missing")
class ClaudeUsageConvertTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.src = Path(cls.tmp.name) / "claude-usage.pptx"
        _write_claude_like_source(cls.src)
        cls.result = _convert(TEMPLATE, cls.src, Path(cls.tmp.name) / "built")
        cls.prs = Presentation(str(cls.result["output"]))
        cls.plan = cls.result["plan"]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _entry(self, source_index: int) -> dict:
        return next(e for e in self.plan["entries"] if e["source_index"] == source_index)

    def test_cover_uses_presentation_title_not_three_parts(self):
        cover = self.prs.slides[0]
        blob = "\n".join(_iter_text(cover))
        _assert_no_dummy(cover, "cover")
        self.assertNotIn("PRESENTATION TITLE", blob)
        self.assertNotIn("THREE PARTS", blob)
        self.assertNotIn("PEOPLE", blob)
        self.assertIn("Claude Usage", blob)
        self.assertNotIn("EVERGLADES FOUNDATION", blob)
        self.assertEqual(self._entry(0)["role"], "TITLE")
        self.assertEqual((self._entry(0).get("placements") or [{}])[0].get("text") or "Claude Usage", "Claude Usage")
        title_hashes = _template_content_photo_hashes(0)
        cover_hashes = _raster_hashes(cover, int(self.prs.slide_width), int(self.prs.slide_height))
        self.assertTrue(
            title_hashes & cover_hashes,
            "cover with no source photo must keep the template title-slide aerial",
        )

    def test_overview_is_not_people_portraits(self):
        overview = next(s for s in self.prs.slides if "Overview" in "\n".join(_iter_text(s)))
        blob = "\n".join(_iter_text(overview))
        _assert_no_dummy(overview, "overview")
        self.assertNotIn("PEOPLE", blob)
        self.assertNotIn("Full name", blob)
        self.assertNotIn("SLIDE TITLE HERE", blob)
        self.assertIn("13 / 19", blob)
        self.assertIn("Time Saved", blob)
        logos = set(self.result["tokens"].get("logo_hashes") or [])
        sample = set(self.result["tokens"].get("sample_photo_hashes") or [])
        for pic in _picture_shapes(overview):
            digest = blob_hash(pic.image.blob)
            if digest in logos or digest in sample:
                continue
            self.assertLess(
                int(pic.width or 0),
                int(1.5 * EMU),
                "256px icons must not occupy 4-inch portrait wells",
            )
            try:
                pic_blob = pic.image.blob
                with Image.open(io.BytesIO(pic_blob)) as im:
                    px = max(im.size)
            except Exception:
                px = 9999
            self.assertGreater(
                px,
                256,
                "UI icons must not be placed as pictures on Overview Sequence/cards",
            )
        texts = [
            s
            for s in overview.shapes
            if getattr(s, "has_text_frame", False)
            and (s.text_frame.text or "").strip()
            and not is_layout_chrome((s.text_frame.text or "").strip())
        ]
        for i, a in enumerate(texts):
            for b in texts[i + 1 :]:
                self.assertLess(
                    max(_overlap(a, b), _overlap(b, a)),
                    0.12,
                    f"Overview text layers overlap {a.text_frame.text[:40]!r} vs {b.text_frame.text[:40]!r}",
                )

    def test_tables_and_charts_do_not_cover_titles(self):
        for slide in self.prs.slides:
            titles = []
            tables = []
            charts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = (shape.text_frame.text or "").strip()
                    if text in {"PRESENTATION TITLE", "CONTENTS", "THREE PARTS", "TWO POINTS", "KICKER", "COMPARISON", "THE POINT", "BY THE NUMBERS"}:
                        continue
                    if text and int(shape.top or 0) < int(2.6 * EMU) and int(shape.height or 0) < int(1.4 * EMU):
                        titles.append(shape)
                if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                    tables.append(shape)
                if shape.shape_type == MSO_SHAPE_TYPE.CHART:
                    charts.append(shape)
            for title in titles:
                for other in tables + charts:
                    self.assertLess(
                        _overlap(title, other),
                        0.15,
                        f"data covers title {title.text_frame.text[:40]!r}",
                    )

    def test_empty_callout_chrome_not_stacked(self):
        overlay = []
        for i, slide in enumerate(self.prs.slides):
            _assert_no_dummy(slide, f"slide {i}")
            for shape in slide.shapes:
                if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
                    continue
                text = (shape.text_frame.text or "").strip() if getattr(shape, "has_text_frame", False) else ""
                if text:
                    continue
                name = (shape.name or "").lower()
                if name.startswith("rectangle") and int(shape.height or 0) > int(1.5 * EMU):
                    try:
                        fill = shape.fill.fore_color.rgb
                        hex_value = f"{fill[0]:02X}{fill[1]:02X}{fill[2]:02X}"
                    except Exception:
                        continue
                    if hex_value in {"F7F9EA", "E8EDC8"}:
                        overlay.append((i, shape.name, hex_value))
        self.assertFalse(overlay, f"source card callouts must not overlay the rebuild: {overlay}")

    def test_three_takeaways_keep_copy_in_cards(self):
        take = next(s for s in self.prs.slides if "Key Takeaways" in "\n".join(_iter_text(s)))
        blob = "\n".join(_iter_text(take))
        self.assertNotIn("SLIDE TITLE HERE", blob)
        self.assertNotIn("Section Title", blob)
        self.assertIn("Adoption is growing quickly", blob)
        self.assertIn("The Science team is the primary adopter", blob)
        self.assertNotIn("Concept one", blob)
        texts = [
            s
            for s in take.shapes
            if getattr(s, "has_text_frame", False)
            and (s.text_frame.text or "").strip()
            and not is_layout_chrome((s.text_frame.text or "").strip())
            and (s.text_frame.text or "").strip() != "Key Takeaways"
        ]
        for i, a in enumerate(texts):
            for b in texts[i + 1 :]:
                self.assertLess(
                    max(_overlap(a, b), _overlap(b, a)),
                    0.12,
                    f"Takeaways stacked layers {a.text_frame.text[:50]!r} vs {b.text_frame.text[:50]!r}",
                )

    def test_no_stacked_content_text_on_any_slide(self):
        for i, slide in enumerate(self.prs.slides):
            texts = [
                s
                for s in slide.shapes
                if getattr(s, "has_text_frame", False)
                and (s.text_frame.text or "").strip()
                and not is_layout_chrome((s.text_frame.text or "").strip())
            ]
            for a_i, a in enumerate(texts):
                for b in texts[a_i + 1 :]:
                    self.assertLess(
                        max(_overlap(a, b), _overlap(b, a)),
                        0.12,
                        f"slide {i + 1} overlapping copy {a.text_frame.text[:40]!r} vs {b.text_frame.text[:40]!r}",
                    )

    def test_who_is_using_splits_tables_off_three_parts(self):
        who_entries = [e for e in self.plan["entries"] if e["source_index"] == 2]
        self.assertGreaterEqual(len(who_entries), 2)
        text_names = " ".join((e.get("template_layout_name") or "") for e in who_entries)
        self.assertTrue(
            any(
                token in text_names
                for token in ("Three Parts", "Comparison", "The Point", "Statement", "Sequence", "Photograph")
            ),
            text_names,
        )

    def test_interior_unused_wells_keep_stock_photos(self):
        stock_by_index = {
            int(proto["index"]): set(proto.get("sample_photo_hashes") or [])
            for proto in self.result["tokens"].get("prototypes") or []
            if proto.get("sample_photo_hashes") and proto.get("role") == "PICTURE"
        }
        self.assertTrue(stock_by_index, "photo sample slides must expose stock photographs")
        slide_w = int(self.prs.slide_width)
        slide_h = int(self.prs.slide_height)
        checked = False
        for entry, slide in zip(self.plan["entries"], self.prs.slides):
            expected = stock_by_index.get(int(entry.get("template_layout_index") or -1))
            if not expected or int(entry.get("source_index") or 0) == 0:
                continue
            checked = True
            self.assertTrue(
                expected & _raster_hashes(slide, slide_w, slide_h),
                "an interior photo layout with no source photograph must keep its template stock",
            )
        if not checked:
            self.assertTrue(stock_by_index)


if __name__ == "__main__":
    unittest.main(verbosity=2)
