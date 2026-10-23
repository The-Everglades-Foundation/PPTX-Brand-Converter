#!/usr/bin/env python3
"""Regression: leftover prototype dummy must not stack on source copy."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from colorways import (  # noqa: E402
    _heading_text_fits,
    is_layout_chrome,
    prototype_text_slots,
)
from helpers import blob_hash  # noqa: E402
from test_margins_images import (  # noqa: E402
    KNIGHT,
    SAMPLE,
    TEMPLATE,
    _convert,
    _raster_hashes,
    _template_content_photo_hashes,
)

EMU_PER_INCH = 914400

LEFTOVER_DUMMY = (
    "Section title goes here",
    "A sentence framing what this section sets out to do",
    "Title template",
    "Bullet one",
    "Bullet two",
    "Description text here",
    "PRESENTER NAME",
    "A subtitle line goes here",
    "Restoring America’s Everglades",
    "Restoring America's Everglades",
    "Two or three sentences of body copy",
    "A second short paragraph",
    "Concept one",
    "Concept two",
    "Concept three",
    "One sentence that carries the whole argument",
    "A short line of support underneath",
    "LABEL HERE",
    "One line explaining what the figure counts",
    "COLUMN ONE",
    "Gallery template",
    "One line naming what this photograph shows",
    "Section title here",
    "One line of description for this section",
)


def _iter_text(slide) -> list[str]:
    found: list[str] = []
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            continue
        if not getattr(shape, "has_text_frame", False):
            continue
        text = (shape.text_frame.text or "").strip()
        if text:
            found.append(text)
    return found


def _assert_no_dummy(slide, label: str) -> None:
    blob = "\n".join(_iter_text(slide))
    for dummy in LEFTOVER_DUMMY:
        if dummy in blob:
            raise AssertionError(f"{label} still has prototype copy {dummy!r}")


class ChromeClassifyTests(unittest.TestCase):
    def test_keeps_layout_labels_and_index_badges(self):
        self.assertTrue(is_layout_chrome("TWO POINTS"))
        self.assertTrue(is_layout_chrome("KICKER"))
        self.assertTrue(is_layout_chrome("01"))
        self.assertTrue(is_layout_chrome("Mac Stone"))
        self.assertFalse(is_layout_chrome("PRESENTATION TITLE"))
        self.assertFalse(is_layout_chrome("SLIDE TITLE HERE"))
        self.assertFalse(is_layout_chrome("Section title goes here"))
        self.assertFalse(is_layout_chrome("Bullet one"))
        self.assertFalse(is_layout_chrome("PRESENTER NAME · MONTH 2026"))


@unittest.skipUnless(TEMPLATE.exists() and SAMPLE.exists(), "template or sample fixture missing")
class SampleOverlapTests(unittest.TestCase):
    def test_sample_has_no_stacked_prototype_copy(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = _convert(TEMPLATE, SAMPLE, Path(tmp))
            prs = Presentation(str(result["output"]))
            for i, slide in enumerate(prs.slides):
                _assert_no_dummy(slide, f"sample slide {i}")


@unittest.skipUnless(TEMPLATE.exists() and KNIGHT.exists(), "Knight fixture not in this environment")
class KnightOverlapTests(unittest.TestCase):
    def test_knight_dummy_cleared_and_source_in_wells(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = _convert(TEMPLATE, KNIGHT, Path(tmp))
            prs = Presentation(str(result["output"]))
            self.assertEqual(len(prs.slides), 6)

            cover = "\n".join(_iter_text(prs.slides[0]))
            _assert_no_dummy(prs.slides[0], "cover")
            self.assertIn("Knight Foundation Grant", cover)
            self.assertNotIn("PRESENTATION TITLE", cover)

            section = prs.slides[3]
            texts = _iter_text(section)
            blob = "\n".join(texts)
            _assert_no_dummy(section, "section divider")
            self.assertIn("Expanded Email Summary and Knight Foundation Profile", blob)

            two = prs.slides[2]
            two_blob = "\n".join(_iter_text(two))
            _assert_no_dummy(two, "two points")
            self.assertIn("Grant Status, Financial Overview, and Key Initiatives", two_blob)
            self.assertIn("Grant Progress and Compliance", two_blob)
            for shape in two.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                text = (shape.text_frame.text or "").strip()
                if text in {"Grant Progress and Compliance", "Key Funded Initiatives"}:
                    self.assertGreater(
                        int(shape.width or 0),
                        int(2.0 * EMU_PER_INCH),
                        "column headings must sit in heading wells, not the 01/02 badges",
                    )

            kicker = "\n".join(_iter_text(prs.slides[4]))
            _assert_no_dummy(prs.slides[4], "kicker")
            self.assertIn("Summary of Jodi Farrell", kicker)
            self.assertIn("Central Coordination Role", kicker)

            profile = "\n".join(_iter_text(prs.slides[5]))
            _assert_no_dummy(prs.slides[5], "profile")
            self.assertIn("Mission and Legacy", profile)

            # Unlocked source photos still present (PR #4).
            source_hashes = {
                shape["image"]["hash"]
                for slide in result["manifest"]["slides"]
                for shape in slide["shapes"]
                if shape.get("kind") == "picture" and not shape.get("drop_logo")
            }
            found = set()
            A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.shape_type not in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
                        continue
                    digest = blob_hash(shape.image.blob)
                    if digest in source_hashes:
                        found.add(digest)
                        self.assertFalse(shape._element.findall(f".//{{{A_NS}}}picLocks"))
            self.assertEqual(found, source_hashes)


def _overlap(a, b) -> float:
    aw, ah = max(1, int(a.width or 0)), max(1, int(a.height or 0))
    left = max(int(a.left or 0), int(b.left or 0))
    top = max(int(a.top or 0), int(b.top or 0))
    right = min(int(a.left or 0) + aw, int(b.left or 0) + int(b.width or 0))
    bottom = min(int(a.top or 0) + ah, int(b.top or 0) + int(b.height or 0))
    if right <= left or bottom <= top:
        return 0.0
    return ((right - left) * (bottom - top)) / (aw * ah)


def _write_two_points_results_source(path: Path) -> None:
    """Claude Usage Results slide as two numbered cards (cover + closer so it is not TITLE/CLOSING)."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    cover = prs.slides.add_slide(blank)
    title = cover.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(10.0), Inches(1.0))
    title.text_frame.paragraphs[0].text = "Claude Usage"

    results = prs.slides.add_slide(blank)
    heading = results.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.2), Inches(0.5))
    heading.text_frame.paragraphs[0].text = "Results & Time Saved"
    left_h = results.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(6.0), Inches(0.8))
    left_h.text_frame.paragraphs[0].text = "Verified output counts × minutes-per-task assumptions."
    left_b = results.shapes.add_textbox(Inches(0.5), Inches(2.6), Inches(6.0), Inches(0.6))
    left_b.text_frame.paragraphs[0].text = "~20 HOURS SAVED"
    right_h = results.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.8), Inches(0.8))
    right_h.text_frame.paragraphs[0].text = "~1,217 minutes across the last 60 days"
    right_b = results.shapes.add_textbox(Inches(7.0), Inches(2.6), Inches(5.8), Inches(1.8))
    right_b.text_frame.paragraphs[0].text = (
        "Conversations account for the majority of time saved (93%), "
        "reflecting the value of quick, iterative Q&A and drafting support alongside document-level work."
    )

    closer = prs.slides.add_slide(blank)
    take = closer.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(10.0), Inches(1.0))
    take.text_frame.paragraphs[0].text = "Key Takeaways"
    prs.save(str(path))


@unittest.skipUnless(TEMPLATE.exists(), "official template missing")
class TwoPointsHeadingFitTests(unittest.TestCase):
    def test_short_titles_fit_heading_wells_long_sentences_do_not(self):
        prs = Presentation(str(TEMPLATE))
        heading = next(
            slot
            for slide in prs.slides
            for slot in prototype_text_slots(slide)
            if (slot.get("text") or "").strip() == "Section Title"
        )
        self.assertTrue(_heading_text_fits("CONNECTORS", heading))
        self.assertFalse(_heading_text_fits("Grant Progress and Compliance", heading))
        self.assertFalse(
            _heading_text_fits("Verified output counts × minutes-per-task assumptions.", heading),
            "a wrapping sentence must not occupy the section-title well",
        )
        self.assertFalse(_heading_text_fits("~1,217 minutes across the last 60 days", heading))


@unittest.skipUnless(TEMPLATE.exists(), "official template missing")
class TwoPointsResultsOverlapTests(unittest.TestCase):
    def test_results_two_points_has_one_layer_per_card(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "results-two-points.pptx"
            _write_two_points_results_source(src)
            result = _convert(TEMPLATE, src, Path(tmp) / "built")
            prs = Presentation(str(result["output"]))
            entry = next(e for e in result["plan"]["entries"] if e["source_index"] == 1)
            self.assertIn(entry.get("template_layout_name"), {"Two Points", "Three Parts", "Sequence"})
            slide = next(s for s in prs.slides if "Results & Time Saved" in "\n".join(_iter_text(s)))
            blob = "\n".join(_iter_text(slide))
            _assert_no_dummy(slide, "two points results")
            self.assertIn("Results & Time Saved", blob)
            self.assertIn("Verified output counts", blob)
            self.assertIn("HOURS SAVED", blob)
            self.assertIn("1,217", blob)
            self.assertIn("Conversations account", blob)

            leftover_labels = [
                t.strip()
                for t in _iter_text(slide)
                if t.strip().upper() in {"HOURS SAVED", "LABEL HERE"} or t.strip() == "HOURS SAVED"
            ]
            self.assertFalse(
                leftover_labels,
                f"prototype-style leftover labels still on the card: {leftover_labels}",
            )

            content = []
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                text = (shape.text_frame.text or "").strip()
                if not text or is_layout_chrome(text) or text == "Results & Time Saved":
                    continue
                content.append(shape)
            for i, a in enumerate(content):
                for b in content[i + 1 :]:
                    self.assertLess(
                        max(_overlap(a, b), _overlap(b, a)),
                        0.12,
                        f"two text frames occupy the same card {a.text_frame.text[:40]!r} vs {b.text_frame.text[:40]!r}",
                    )

            columns: list[list] = []
            for shape in sorted(content, key=lambda s: (int(s.left or 0), int(s.top or 0))):
                placed = False
                center = int(shape.left or 0) + int(shape.width or 0) / 2
                for col in columns:
                    col_left = min(int(s.left or 0) for s in col)
                    col_right = max(int(s.left or 0) + int(s.width or 0) for s in col)
                    if col_left - int(0.4 * EMU_PER_INCH) <= center <= col_right + int(0.4 * EMU_PER_INCH):
                        col.append(shape)
                        placed = True
                        break
                if not placed:
                    columns.append([shape])
            self.assertGreaterEqual(len(content), 2, "both result points should be on the slide")
            stock = _template_content_photo_hashes(int(entry.get("template_layout_index") or 0))
            if stock:
                found = _raster_hashes(slide, int(prs.slide_width), int(prs.slide_height))
                self.assertTrue(found & stock, "unused photo-well stock must stay")


if __name__ == "__main__":
    unittest.main(verbosity=2)
