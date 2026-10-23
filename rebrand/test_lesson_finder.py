#!/usr/bin/env python3
"""Regression: Everglades Lesson Finder-style decks convert cleanly (blue)."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
sys.path.insert(0, str(ROOT))

from build import build_presentation  # noqa: E402
from colorways import is_layout_chrome  # noqa: E402
from helpers import blob_hash, dump_json, is_decorative_icon, is_org_chrome  # noqa: E402
from inspect_source import inspect_source  # noqa: E402
from inspect_template import inspect_template  # noqa: E402
from plan import build_plan  # noqa: E402
from test_margins_images import TEMPLATE, _raster_hashes, _template_content_photo_hashes  # noqa: E402
from test_overlap_layers import _assert_no_dummy, _iter_text  # noqa: E402

EMU = 914400
REAL_DECK = Path(
    "/cursor/stores/bc-d2ce698c-eb2f-4f8d-a4a5-1d14dc2ece8e/media/lesson-finder-original.pptx"
)
FIXTURE = REPO / "fixtures" / "lesson_finder.pptx"


def _icon_png() -> bytes:
    im = Image.new("RGBA", (256, 256), (0, 0, 0, 0))
    draw = ImageDraw.Draw(im)
    draw.ellipse((32, 32, 224, 224), fill=(46, 120, 51, 255))
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _grass_png() -> bytes:
    im = Image.new("RGBA", (2600, 110), (0, 0, 0, 0))
    draw = ImageDraw.Draw(im)
    draw.rectangle((0, 40, 2600, 110), fill=(130, 160, 70, 255))
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _screenshot_jpeg() -> bytes:
    im = Image.new("RGB", (640, 360), (240, 248, 252))
    draw = ImageDraw.Draw(im)
    draw.rectangle((20, 20, 620, 80), fill=(0, 172, 191))
    draw.rectangle((40, 110, 300, 320), fill=(255, 255, 255))
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=80)
    return buf.getvalue()


def write_lesson_finder_source(path: Path) -> None:
    """Compact Lesson Finder analogue: org kicker, 256px icons, two points, steps, chart."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    icon = _icon_png()
    grass = _grass_png()
    shot = _screenshot_jpeg()

    cover = prs.slides.add_slide(blank)
    kicker = cover.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.0), Inches(0.4))
    kicker.text_frame.paragraphs[0].text = "THE EVERGLADES FOUNDATION · EVERGLADES LITERACY PROGRAM"
    title = cover.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(9.0), Inches(0.8))
    title.text_frame.paragraphs[0].text = "Everglades Lesson Finder"
    body = cover.shapes.add_textbox(Inches(0.5), Inches(2.8), Inches(9.0), Inches(2.2))
    tf = body.text_frame
    tf.paragraphs[0].text = (
        "An AI assistant on evergladesliteracy.org that helps any Florida teacher "
        "find the right Everglades lesson by topic, grade level, or standard."
    )
    cover.shapes.add_picture(io.BytesIO(grass), Inches(8.6), Inches(1.6), Inches(4.0), Inches(0.17))

    why = prs.slides.add_slide(blank)
    wt = why.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.0), Inches(0.6))
    wt.text_frame.paragraphs[0].text = "Why this project"
    left = why.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(7.5), Inches(4.4))
    ltf = left.text_frame
    ltf.paragraphs[0].text = "Today"
    p = ltf.add_paragraph()
    p.text = (
        "The Teacher Toolkit holds 41 lessons across PreK–12, each aligned to standards "
        "and Fundamental Concepts."
    )
    p = ltf.add_paragraph()
    p.text = "With the Lesson Finder"
    p = ltf.add_paragraph()
    p.text = (
        'Teachers ask in plain English, right on our website — '
        '“Find a 5th-grade lesson on invasive species.”'
    )
    why.shapes.add_picture(io.BytesIO(icon), Inches(8.6), Inches(1.6), Inches(4.0), Inches(4.0))

    does = prs.slides.add_slide(blank)
    dt = does.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.0), Inches(0.6))
    dt.text_frame.paragraphs[0].text = "What it does"
    db = does.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(9.0), Inches(4.8))
    dtf = db.text_frame
    lines = [
        "Finds lessons by topic",
        "Titles, grade levels, and a one-line summary.",
        "Maps standards & concepts",
        "Shows the NGSSS codes each lesson supports.",
        "Matches the grade",
        "Suggests the closest grade if none exists.",
        "Hands over the materials",
        "Every answer ends with direct download links.",
    ]
    dtf.paragraphs[0].text = lines[0]
    for line in lines[1:]:
        dtf.add_paragraph().text = line
    does.shapes.add_picture(io.BytesIO(icon), Inches(8.6), Inches(1.6), Inches(4.0), Inches(4.0))

    pic = prs.slides.add_slide(blank)
    pic.shapes.add_picture(io.BytesIO(shot), Inches(1.96), Inches(0.67), Inches(6.0), Inches(4.5))

    two = prs.slides.add_slide(blank)
    tt = two.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.0), Inches(0.6))
    tt.text_frame.paragraphs[0].text = "What teachers will see"
    col_a = two.shapes.add_textbox(Inches(0.5), Inches(1.75), Inches(5.8), Inches(4.5))
    col_a.text_frame.paragraphs[0].text = (
        "A teacher opens the widget on the Teacher Toolkit page — tap a suggestion or ask anything."
    )
    col_b = two.shapes.add_textbox(Inches(6.8), Inches(1.75), Inches(5.8), Inches(4.5))
    col_b.text_frame.paragraphs[0].text = (
        "One answer: the lesson, grade, standard, concept — and the downloads, ready to click."
    )

    costs = prs.slides.add_slide(blank)
    ct = costs.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.0), Inches(0.6))
    ct.text_frame.paragraphs[0].text = "What it costs"
    data = CategoryChartData()
    data.categories = ["Build", "Host", "Support"]
    data.add_series("USD", (2400, 600, 300))
    costs.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.55), Inches(11.9), Inches(5.3), data)

    steps = prs.slides.add_slide(blank)
    st = steps.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.0), Inches(0.6))
    st.text_frame.paragraphs[0].text = "Next steps"
    left_s = steps.shapes.add_textbox(Inches(0.5), Inches(1.75), Inches(5.8), Inches(4.5))
    ltf = left_s.text_frame
    for i, line in enumerate(("Step 1", "Build the index", "Education team completes the spreadsheet.", "Step 2", "Deploy & test internally", "Service goes live on a private link.")):
        if i == 0:
            ltf.paragraphs[0].text = line
        else:
            ltf.add_paragraph().text = line
    right_s = steps.shapes.add_textbox(Inches(6.8), Inches(1.75), Inches(5.8), Inches(4.5))
    rtf = right_s.text_frame
    for i, line in enumerate(("Step 3", "Teacher pilot", "Champion School teachers try it.", "Step 4", "Public launch", "Widget embedded on the Teacher Toolkit page.")):
        if i == 0:
            rtf.paragraphs[0].text = line
        else:
            rtf.add_paragraph().text = line

    prs.save(str(path))


def _convert_blue(template: Path, source: Path, out_dir: Path) -> dict:
    tokens = inspect_template(template, out_dir, colorway="blue")
    manifest = inspect_source(source, out_dir)
    plan = build_plan(tokens, manifest, colorway="blue")
    dump_json(out_dir / "plan.json", plan)
    meta = build_presentation(
        template_path=template,
        source_path=source,
        out_dir=out_dir,
        tokens=tokens,
        manifest=manifest,
        plan=plan,
        brand_md=None,
        colorway="blue",
    )
    return {"tokens": tokens, "manifest": manifest, "plan": plan, "meta": meta, "output": Path(meta["output"])}


def _overlap(a, b) -> float:
    aw, ah = max(1, int(a.width or 0)), max(1, int(a.height or 0))
    left = max(int(a.left or 0), int(b.left or 0))
    top = max(int(a.top or 0), int(b.top or 0))
    right = min(int(a.left or 0) + aw, int(b.left or 0) + int(b.width or 0))
    bottom = min(int(a.top or 0) + ah, int(b.top or 0) + int(b.height or 0))
    if right <= left or bottom <= top:
        return 0.0
    return ((right - left) * (bottom - top)) / (aw * ah)


def _slide_blob(slide) -> str:
    return "\n".join(_iter_text(slide))


class HelperTests(unittest.TestCase):
    def test_org_kicker_is_chrome_long_body_is_not(self):
        self.assertTrue(is_org_chrome("THE EVERGLADES FOUNDATION · EVERGLADES LITERACY PROGRAM"))
        self.assertTrue(is_org_chrome("The Everglades Foundation"))
        self.assertFalse(
            is_org_chrome(
                "THE EVERGLADES FOUNDATION · EVERGLADES LITERACY PROGRAM "
                "Everglades Lesson Finder An AI assistant on evergladesliteracy.org"
            )
        )

    def test_grass_strip_and_256_icon_are_decorative(self):
        self.assertTrue(
            is_decorative_icon({"image": {"px": [2600, 110], "bytes": 13431, "ext": "png"}, "area_ratio": 0.007})
        )
        self.assertTrue(
            is_decorative_icon({"image": {"px": [256, 256], "bytes": 5402, "ext": "png"}, "area_ratio": 0.16})
        )
        self.assertFalse(
            is_decorative_icon({"image": {"px": [640, 360], "bytes": 40000, "ext": "jpeg"}, "area_ratio": 0.27})
        )


@unittest.skipUnless(TEMPLATE.exists(), "official template missing")
class LessonFinderConvertTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.src = Path(cls.tmp.name) / "lesson-finder.pptx"
        write_lesson_finder_source(cls.src)
        FIXTURE.parent.mkdir(parents=True, exist_ok=True)
        if not FIXTURE.exists() or FIXTURE.stat().st_size < 1000:
            write_lesson_finder_source(FIXTURE)
        cls.result = _convert_blue(TEMPLATE, cls.src, Path(cls.tmp.name) / "built")
        cls.prs = Presentation(str(cls.result["output"]))
        cls.plan = cls.result["plan"]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _by_title(self, needle: str):
        for slide in self.prs.slides:
            if needle in _slide_blob(slide):
                return slide
        self.fail(f"no slide contains {needle!r}")

    def test_colorway_is_blue(self):
        self.assertEqual(self.plan.get("colorway"), "blue")

    def test_cover_uses_lesson_finder_title_and_keeps_aerial(self):
        cover = self.prs.slides[0]
        blob = _slide_blob(cover)
        _assert_no_dummy(cover, "cover")
        self.assertIn("Everglades Lesson Finder", blob)
        self.assertNotIn("PRESENTATION TITLE", blob)
        self.assertNotIn("THE EVERGLADES FOUNDATION · EVERGLADES LITERACY PROGRAM", blob)
        self.assertEqual(self.plan["entries"][0]["role"], "TITLE")
        self.assertEqual(self.plan["entries"][0].get("template_layout_name"), "Presentation Title")
        aerial = _template_content_photo_hashes(0)
        cover_hashes = _raster_hashes(cover, int(self.prs.slide_width), int(self.prs.slide_height))
        self.assertTrue(aerial & cover_hashes, "cover must keep the unused title-slide aerial")
        grass = blob_hash(_grass_png())
        self.assertNotIn(grass, cover_hashes, "grass divider must not replace the cover photo")

    def test_two_points_today_vs_finder_expands_and_keeps_stock(self):
        slide = self._by_title("Why this project")
        blob = _slide_blob(slide)
        _assert_no_dummy(slide, "why")
        self.assertNotIn("SLIDE TITLE HERE", blob)
        self.assertIn("Today", blob)
        self.assertIn("With the Lesson Finder", blob)
        bodies = [
            s
            for s in slide.shapes
            if getattr(s, "has_text_frame", False)
            and "Teacher Toolkit" in (s.text_frame.text or "")
        ]
        self.assertTrue(bodies)
        self.assertGreater(
            int(bodies[0].height or 0),
            int(2.0 * EMU),
            "Two Points body well must grow so the Today copy is not clipped",
        )
        layout_index = next(
            int(entry.get("template_layout_index") or -1)
            for entry in self.plan["entries"]
            if "Why this project" in " ".join(
                (p.get("text") or "") for p in entry.get("placements") or []
            )
        )
        stock = _template_content_photo_hashes(layout_index) if layout_index >= 0 else set()
        found = _raster_hashes(slide, int(self.prs.slide_width), int(self.prs.slide_height))
        if stock:
            self.assertTrue(found & stock, "unused photo-well stock must stay when the source has no photograph")
        self.assertNotIn(blob_hash(_icon_png()), found, "256px icons must not occupy photo wells")

    def test_capabilities_and_steps_use_sequence_not_crammed_two_points(self):
        does = self._by_title("What it does")
        self.assertIn("Finds lessons by topic", _slide_blob(does))
        self.assertNotIn("SLIDE TITLE HERE", _slide_blob(does))
        _assert_no_dummy(does, "what it does")

        deck = "\n".join(_slide_blob(slide) for slide in self.prs.slides)
        self.assertNotIn("SLIDE TITLE HERE", deck)
        self.assertIn("Step 1", deck)
        self.assertIn("Step 2", deck)
        self.assertIn("Step 3", deck)
        self.assertIn("Step 4", deck)

    def test_chart_stays_on_one_slide_without_covering_title(self):
        cost_entries = [e for e in self.plan["entries"] if e["source_index"] == 5]
        self.assertEqual(len(cost_entries), 1, "title+chart must not split onto a (cont.) slide")
        slide = self._by_title("What it costs")
        self.assertNotIn("(cont.)", _slide_blob(slide))
        titles = []
        charts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = (shape.text_frame.text or "").strip()
                if text == "What it costs":
                    titles.append(shape)
            if shape.shape_type == MSO_SHAPE_TYPE.CHART:
                charts.append(shape)
        self.assertTrue(charts)
        self.assertTrue(titles)
        for title in titles:
            for chart in charts:
                self.assertLess(_overlap(title, chart), 0.08, "chart overlaps What it costs")

    def test_screenshot_photo_is_kept_and_no_dummy(self):
        shot_hash = blob_hash(_screenshot_jpeg())
        found = False
        for slide in self.prs.slides:
            _assert_no_dummy(slide, "screenshot-ish")
            for pic in slide.shapes:
                if pic.shape_type not in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
                    continue
                if blob_hash(pic.image.blob) == shot_hash:
                    found = True
        self.assertTrue(found, "source screenshot must be placed")

    def test_no_stacked_content_text(self):
        for i, slide in enumerate(self.prs.slides):
            texts = [
                s
                for s in slide.shapes
                if getattr(s, "has_text_frame", False)
                and (s.text_frame.text or "").strip()
                and not is_layout_chrome((s.text_frame.text or "").strip())
            ]
            for a_i, a in enumerate(texts):
                for b in texts[a_i + 1 :]:
                    self.assertLess(
                        max(_overlap(a, b), _overlap(b, a)),
                        0.12,
                        f"slide {i + 1} overlapping copy {a.text_frame.text[:40]!r} vs {b.text_frame.text[:40]!r}",
                    )


@unittest.skipUnless(TEMPLATE.exists() and REAL_DECK.exists(), "real Lesson Finder deck not in this environment")
class RealLessonFinderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.result = _convert_blue(TEMPLATE, REAL_DECK, Path(cls.tmp.name) / "built")
        cls.prs = Presentation(str(cls.result["output"]))
        cls.plan = cls.result["plan"]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_real_deck_cover_and_no_split_chart(self):
        self.assertEqual(self.plan.get("colorway"), "blue")
        self.assertEqual(len(self.prs.slides), 10)
        cover = "\n".join(_iter_text(self.prs.slides[0]))
        self.assertIn("Everglades Lesson Finder", cover)
        self.assertNotIn("THE EVERGLADES FOUNDATION · EVERGLADES LITERACY PROGRAM", cover)
        cost = [e for e in self.plan["entries"] if "What it costs" in str(e)]
        self.assertTrue(any(e.get("template_layout_name") == "The Point" for e in self.plan["entries"]))
        self.assertFalse(any("(cont.)" in "\n".join(_iter_text(s)) for s in self.prs.slides))
        why = next(s for s in self.prs.slides if "Why this project" in "\n".join(_iter_text(s)))
        self.assertIn("Today", "\n".join(_iter_text(why)))
        self.assertIn("With the Lesson Finder", "\n".join(_iter_text(why)))
        for slide in self.prs.slides:
            _assert_no_dummy(slide, "real lesson finder")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--write-fixture":
        FIXTURE.parent.mkdir(parents=True, exist_ok=True)
        write_lesson_finder_source(FIXTURE)
        print(f"Wrote {FIXTURE} ({FIXTURE.stat().st_size} bytes)")
        raise SystemExit(0)
    unittest.main(verbosity=2)
