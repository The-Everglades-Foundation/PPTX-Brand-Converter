#!/usr/bin/env python3
"""Regression: ForEverglades / Knight-style green convert keeps titles, pairs KPIs."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
sys.path.insert(0, str(ROOT))

from helpers import (  # noqa: E402
    blob_hash,
    is_metric_line,
    is_title_like,
    looks_like_metric_pairs,
    looks_like_photo_caption,
    pair_metric_label_items,
    title_candidate_score,
)
from inspect_source import pick_title  # noqa: E402
from test_margins_images import KNIGHT, TEMPLATE, _convert  # noqa: E402
from test_overlap_layers import _assert_no_dummy, _iter_text  # noqa: E402
from colorways import is_layout_chrome  # noqa: E402

EMU = 914400
FIXTURE = REPO / "fixtures" / "knight_foreverglades.pptx"


def _jpeg(color: tuple[int, int, int], size: tuple[int, int] = (640, 420)) -> bytes:
    im = Image.new("RGB", size, color)
    draw = ImageDraw.Draw(im)
    draw.rectangle((20, 20, size[0] - 20, size[1] - 20), outline=(255, 255, 255), width=6)
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=70)
    return buf.getvalue()


def _wide_banner_jpeg() -> bytes:
    return _jpeg((70, 110, 60), (1568, 532))


def _hero_jpeg() -> bytes:
    return _jpeg((30, 80, 90), (1372, 896))


def _right_jpeg() -> bytes:
    return _jpeg((120, 90, 40), (1372, 896))


def _banner_photo_jpeg() -> bytes:
    return _jpeg((80, 50, 70), (1344, 896))


def write_knight_foreverglades_source(path: Path) -> None:
    """Compact analogue of the ForEverglades Knight green convert."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    hero = _hero_jpeg()
    wide = _wide_banner_jpeg()
    right = _right_jpeg()
    banner = _banner_photo_jpeg()

    photo_only = prs.slides.add_slide(blank)
    photo_only.shapes.add_picture(io.BytesIO(wide), Inches(0), Inches(1.6), Inches(13.333), Inches(4.5))

    intro = prs.slides.add_slide(blank)
    intro.shapes.add_picture(io.BytesIO(hero), Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    title = intro.shapes.add_textbox(Inches(0.6), Inches(5.8), Inches(8.0), Inches(0.7))
    title.text_frame.paragraphs[0].text = "INTRODUCING FOREVERGLADES"

    film = prs.slides.add_slide(blank)
    ft = film.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(12.0), Inches(0.7))
    ft.text_frame.paragraphs[0].text = "ONE FILM, A LIBRARY"
    fb = film.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(11.0), Inches(3.8))
    fb.text_frame.paragraphs[0].text = (
        "The Everglades Foundation will own the film and build it into a scalable, "
        "multi-year educational platform. From it we will create short-form video "
        "segments for classrooms and the public."
    )

    metrics = prs.slides.add_slide(blank)
    mt = metrics.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.6))
    mt.text_frame.paragraphs[0].text = "1 MILLION"
    ms = metrics.shapes.add_textbox(Inches(0.6), Inches(0.95), Inches(12.0), Inches(0.4))
    ms.text_frame.paragraphs[0].text = "potential viewers in year one"
    pairs = [
        ("312,000", "students reached to date"),
        ("~10,000", "FL teachers (2,400+ Miami-Dade)"),
    ]
    for i, (num, label) in enumerate(pairs):
        x = 0.6 + i * 6.2
        nbox = metrics.shapes.add_textbox(Inches(x), Inches(2.2), Inches(5.6), Inches(0.6))
        nbox.text_frame.paragraphs[0].text = num
        lbox = metrics.shapes.add_textbox(Inches(x), Inches(2.9), Inches(5.6), Inches(1.2))
        lbox.text_frame.paragraphs[0].text = label

    digital = prs.slides.add_slide(blank)
    dt = digital.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(7.4), Inches(0.55))
    dt.text_frame.paragraphs[0].text = "DIGITAL GROUNDWORK"
    db = digital.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(7.4), Inches(5.6))
    dtf = db.text_frame
    dtf.paragraphs[0].text = "UNIFIED DATA"
    p = dtf.add_paragraph()
    p.text = (
        "Every program now runs on one integrated Salesforce system, streamlining "
        "events, donor relationships, and outreach."
    )
    p = dtf.add_paragraph()
    p.text = "ETHICAL AI"
    p = dtf.add_paragraph()
    p.text = "Staff are trained to use AI ethically across departments."
    p = dtf.add_paragraph()
    p.text = "STORYTELLING"
    p = dtf.add_paragraph()
    p.text = "The Restoration Works series launches with featured Miami-Dade projects."
    digital.shapes.add_picture(io.BytesIO(right), Inches(8.4), Inches(0.8), Inches(4.5), Inches(5.8))

    partner = prs.slides.add_slide(blank)
    callout = partner.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.2), Inches(0.85))
    callout.text_frame.paragraphs[0].text = (
        '$350,000: on-screen credit as "Major Production Funding Provided By...", '
        "repeated at the film's conclusion."
    )
    left = partner.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(6.0), Inches(3.2))
    ltf = left.text_frame
    ltf.paragraphs[0].text = "Support from the Knight Foundation will be recognized as follows."
    p = ltf.add_paragraph()
    p.text = "Under $350,000: on-screen credit at the conclusion of the film."
    right_col = partner.shapes.add_textbox(Inches(6.7), Inches(1.4), Inches(6.0), Inches(3.2))
    rtf = right_col.text_frame
    rtf.paragraphs[0].text = "Flexible giving: donations may be pledged over three years to ease the commitment."
    named = partner.shapes.add_textbox(Inches(0.5), Inches(4.8), Inches(6.0), Inches(0.5))
    named.text_frame.paragraphs[0].text = "WAYS TO PARTNER"
    partner.shapes.add_picture(io.BytesIO(banner), Inches(4.2), Inches(5.5), Inches(5.0), Inches(1.8))

    ask = prs.slides.add_slide(blank)
    at = ask.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(7.4), Inches(0.5))
    at.text_frame.paragraphs[0].text = "THE ASK"
    ab = ask.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(7.4), Inches(4.8))
    ab.text_frame.paragraphs[0].text = (
        "We hope the Knight Foundation will amend our current grant to include "
        "support for ForEverglades, creating a timeless resource for youth engagement."
    )
    ask.shapes.add_picture(io.BytesIO(right), Inches(8.4), Inches(0.8), Inches(4.5), Inches(5.8))

    prs.save(str(path))


def _slide_blob(slide) -> str:
    return "\n".join(_iter_text(slide))


def _overlap(a, b) -> float:
    aw, ah = max(1, int(a.width or 0)), max(1, int(a.height or 0))
    left = max(int(a.left or 0), int(b.left or 0))
    top = max(int(a.top or 0), int(b.top or 0))
    right = min(int(a.left or 0) + aw, int(b.left or 0) + int(b.width or 0))
    bottom = min(int(a.top or 0) + ah, int(b.top or 0) + int(b.height or 0))
    if right <= left or bottom <= top:
        return 0.0
    return ((right - left) * (bottom - top)) / (aw * ah)


class HelperTests(unittest.TestCase):
    def test_metric_pairing(self):
        items = ["312,000", "students reached to date", "~10,000", "FL teachers (2,400+ Miami-Dade)"]
        self.assertTrue(looks_like_metric_pairs(items))
        self.assertTrue(is_metric_line("312,000"))
        self.assertTrue(is_metric_line("~10,000"))
        self.assertTrue(is_metric_line("1 MILLION"))
        self.assertEqual(
            pair_metric_label_items(items),
            [["312,000", "students reached to date"], ["~10,000", "FL teachers (2,400+ Miami-Dade)"]],
        )

    def test_title_like_vs_caption_and_money(self):
        self.assertTrue(is_title_like("INTRODUCING FOREVERGLADES"))
        self.assertTrue(is_title_like("WAYS TO PARTNER"))
        self.assertFalse(
            is_title_like(
                '$350,000: on-screen credit as "Major Production Funding Provided By...", '
                "repeated at the film's conclusion."
            )
        )
        self.assertTrue(looks_like_photo_caption("Marsh survey site"))
        self.assertFalse(looks_like_photo_caption("INTRODUCING FOREVERGLADES"))
        self.assertGreater(
            title_candidate_score("WAYS TO PARTNER"),
            title_candidate_score(
                '$350,000: on-screen credit as "Major Production Funding Provided By...", '
                "repeated at the film's conclusion."
            ),
        )

    def test_pick_title_prefers_all_caps_over_money_callout(self):
        emu = 914400
        shapes = [
            {
                "kind": "text",
                "name": "Callout",
                "placeholder_type": None,
                "drop_footer": False,
                "text": '$350,000: on-screen credit as "Major Production Funding Provided By...", repeated at the film\'s conclusion.',
                "bounds": {"top": 0, "left": 0, "width": 10 * emu, "height": emu},
                "text_frame": {
                    "paragraphs": [
                        {
                            "text": '$350,000: on-screen credit as "Major Production Funding Provided By...", repeated at the film\'s conclusion.'
                        }
                    ]
                },
            },
            {
                "kind": "text",
                "name": "Label",
                "placeholder_type": None,
                "drop_footer": False,
                "text": "WAYS TO PARTNER",
                "bounds": {"top": int(4.8 * emu), "left": 0, "width": 6 * emu, "height": int(0.5 * emu)},
                "text_frame": {"paragraphs": [{"text": "WAYS TO PARTNER"}]},
            },
        ]
        title, _sub = pick_title(shapes)
        self.assertEqual(title, "WAYS TO PARTNER")


@unittest.skipUnless(TEMPLATE.exists(), "official template missing")
class KnightForevergladesConvertTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.src = Path(cls.tmp.name) / "knight-foreverglades.pptx"
        write_knight_foreverglades_source(cls.src)
        FIXTURE.parent.mkdir(parents=True, exist_ok=True)
        write_knight_foreverglades_source(FIXTURE)
        cls.result = _convert(TEMPLATE, cls.src, Path(cls.tmp.name) / "built")
        cls.prs = Presentation(str(cls.result["output"]))
        cls.plan = cls.result["plan"]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _by_title(self, needle: str):
        for slide in self.prs.slides:
            if needle in _slide_blob(slide):
                return slide
        self.fail(f"no slide contains {needle!r}")

    def _entry_for(self, needle: str) -> dict:
        for entry in self.plan["entries"]:
            blob = " ".join(
                str(p.get("text") or "")
                + " "
                + " ".join((para.get("text") or "") for para in (p.get("paragraphs") or []))
                for p in entry.get("placements") or []
            )
            if needle in blob:
                return entry
        self.fail(f"no plan entry contains {needle!r}")

    def test_intro_title_is_not_caption_layout(self):
        entry = self._entry_for("INTRODUCING FOREVERGLADES")
        self.assertNotEqual(entry.get("role"), "PICTURE")
        name = (entry.get("template_layout_name") or "").lower()
        self.assertNotIn("caption", name)
        slide = self._by_title("INTRODUCING FOREVERGLADES")
        blob = _slide_blob(slide)
        _assert_no_dummy(slide, "intro")
        self.assertNotIn("CAPTION", blob)
        titles = [
            s
            for s in slide.shapes
            if getattr(s, "has_text_frame", False)
            and "INTRODUCING FOREVERGLADES" in (s.text_frame.text or "")
        ]
        self.assertTrue(titles)
        self.assertGreater(int(titles[0].width or 0), int(6.0 * EMU), "title must use a title well, not a caption strip")

    def test_metric_number_stays_with_its_label(self):
        slide = self._by_title("1 MILLION")
        blob = _slide_blob(slide)
        _assert_no_dummy(slide, "metrics")
        self.assertIn("312,000", blob)
        self.assertIn("students reached to date", blob)
        self.assertIn("~10,000", blob)
        self.assertIn("FL teachers", blob)
        self.assertNotIn("SEQUENCE", blob)
        self.assertIn("TWO POINTS", blob)
        num = next(s for s in slide.shapes if getattr(s, "has_text_frame", False) and "312,000" in (s.text_frame.text or ""))
        label = next(
            s
            for s in slide.shapes
            if getattr(s, "has_text_frame", False) and "students reached to date" in (s.text_frame.text or "")
        )
        num_cx = int(num.left) + int(num.width) / 2
        label_cx = int(label.left) + int(label.width) / 2
        self.assertLess(abs(num_cx - label_cx), int(3.0 * EMU), "312,000 must stay with students reached")
        other = next(s for s in slide.shapes if getattr(s, "has_text_frame", False) and "~10,000" in (s.text_frame.text or ""))
        self.assertGreater(abs(num_cx - (int(other.left) + int(other.width) / 2)), int(3.0 * EMU))

    def test_ways_to_partner_is_the_title(self):
        slide = self._by_title("WAYS TO PARTNER")
        blob = _slide_blob(slide)
        _assert_no_dummy(slide, "partner")
        self.assertIn("Major Production Funding", blob)
        titles = [
            s
            for s in slide.shapes
            if getattr(s, "has_text_frame", False)
            and (s.text_frame.text or "").strip() == "WAYS TO PARTNER"
        ]
        self.assertTrue(titles, "WAYS TO PARTNER must occupy its own title well")
        self.assertGreater(int(titles[0].width or 0), int(8.0 * EMU))
        money = [
            s
            for s in slide.shapes
            if getattr(s, "has_text_frame", False)
            and "$350,000" in (s.text_frame.text or "")
            and "WAYS TO PARTNER" not in (s.text_frame.text or "")
        ]
        self.assertTrue(money, "$350,000 callout stays in the body, not as the slide title")

    def test_digital_groundwork_uses_three_parts(self):
        slide = self._by_title("DIGITAL GROUNDWORK")
        blob = _slide_blob(slide)
        _assert_no_dummy(slide, "digital")
        self.assertIn("THREE PARTS", blob)
        self.assertIn("UNIFIED DATA", blob)
        self.assertIn("ETHICAL AI", blob)
        self.assertIn("STORYTELLING", blob)
        self.assertNotIn("KICKER", blob)

    def test_ask_keeps_title_and_photos_unlocked(self):
        slide = self._by_title("THE ASK")
        blob = _slide_blob(slide)
        _assert_no_dummy(slide, "ask")
        self.assertIn("THE ASK", blob)
        self.assertIn("amend our current grant", blob)
        source_hashes = {
            shape["image"]["hash"]
            for src_slide in self.result["manifest"]["slides"]
            for shape in src_slide["shapes"]
            if shape.get("kind") == "picture" and not shape.get("drop_logo") and not shape.get("drop_icon")
        }
        found = set()
        for out in self.prs.slides:
            for pic in out.shapes:
                if pic.shape_type not in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
                    continue
                digest = blob_hash(pic.image.blob)
                if digest in source_hashes:
                    found.add(digest)
                    locks = pic._element.findall(
                        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}picLocks"
                    )
                    self.assertFalse(locks, "source photos must be unlocked")
        self.assertEqual(found, source_hashes)

    def test_no_dummy_or_stacked_copy(self):
        for i, slide in enumerate(self.prs.slides):
            _assert_no_dummy(slide, f"knight foreverglades {i}")
            texts = [
                s
                for s in slide.shapes
                if getattr(s, "has_text_frame", False)
                and (s.text_frame.text or "").strip()
                and not is_layout_chrome((s.text_frame.text or "").strip())
            ]
            for a_i, a in enumerate(texts):
                for b in texts[a_i + 1 :]:
                    self.assertLess(
                        max(_overlap(a, b), _overlap(b, a)),
                        0.12,
                        f"slide {i + 1} overlapping copy {a.text_frame.text[:40]!r} vs {b.text_frame.text[:40]!r}",
                    )


@unittest.skipUnless(TEMPLATE.exists() and KNIGHT.exists(), "Knight store fixture missing")
class ExistingKnightFixtureTests(unittest.TestCase):
    def test_grant_brief_still_converts_cleanly(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = _convert(TEMPLATE, KNIGHT, Path(tmp))
            prs = Presentation(str(result["output"]))
            self.assertEqual(len(prs.slides), 6)
            cover = _slide_blob(prs.slides[0])
            self.assertIn("Knight Foundation Grant", cover)
            self.assertIn("PRESENTATION TITLE", cover)
            _assert_no_dummy(prs.slides[0], "existing knight cover")
            two = next(s for s in prs.slides if "TWO POINTS" in _slide_blob(s))
            self.assertIn("Grant Progress and Compliance", _slide_blob(two))


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--write-fixture":
        FIXTURE.parent.mkdir(parents=True, exist_ok=True)
        write_knight_foreverglades_source(FIXTURE)
        print(f"Wrote {FIXTURE} ({FIXTURE.stat().st_size} bytes)")
        raise SystemExit(0)
    unittest.main(verbosity=2)
