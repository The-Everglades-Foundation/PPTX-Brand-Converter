#!/usr/bin/env python3
"""Regression tests: template margins, unlocked source photos, unused stock art kept."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402
from PIL import Image  # noqa: E402

from colorways import content_photo_frames, sanitize_new_picture  # noqa: E402
from helpers import blob_hash, content_safe_box, is_brand_lockup  # noqa: E402
from inspect_source import inspect_source  # noqa: E402
from inspect_template import inspect_template  # noqa: E402
from plan import build_plan  # noqa: E402
from build import build_presentation  # noqa: E402

TEMPLATE = REPO / "templates" / "everglades.pptx"
SAMPLE = REPO / "fixtures" / "sample.pptx"
KNIGHT = Path(
    "/cursor/stores/bc-d2ce698c-eb2f-4f8d-a4a5-1d14dc2ece8e/internal/"
    "Knight_Foundation_Text_and_Images_No_Branding.pptx"
)
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
EMU_PER_INCH = 914400


def _png_bytes(color: tuple[int, int, int], size: tuple[int, int] = (80, 60)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


def _pic_locks(shape) -> list[dict[str, str]]:
    found = []
    for node in shape._element.findall(f".//{{{A_NS}}}picLocks"):
        found.append(dict(node.attrib))
    return found


def _convert(template: Path, source: Path, out_dir: Path) -> dict:
    tokens = inspect_template(template, out_dir, colorway="green")
    manifest = inspect_source(source, out_dir)
    plan = build_plan(tokens, manifest, colorway="green")
    from helpers import dump_json

    dump_json(out_dir / "plan.json", plan)
    meta = build_presentation(
        template_path=template,
        source_path=source,
        out_dir=out_dir,
        tokens=tokens,
        manifest=manifest,
        plan=plan,
        brand_md=None,
        colorway="green",
    )
    return {"tokens": tokens, "manifest": manifest, "plan": plan, "meta": meta, "output": Path(meta["output"])}


def _picture_shapes(slide):
    return [s for s in slide.shapes if s.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}]


class HelperTests(unittest.TestCase):
    def test_brand_lockup_detects_top_left_mark(self):
        slide_w, slide_h = 20 * EMU_PER_INCH, int(11.25 * EMU_PER_INCH)
        slide_area = slide_w * slide_h
        left, top = int(1.15 * EMU_PER_INCH), int(1.04 * EMU_PER_INCH)
        width, height = int(3.44 * EMU_PER_INCH), int(2.09 * EMU_PER_INCH)
        self.assertTrue(is_brand_lockup(width * height, slide_area, left, top, width, height, slide_w, slide_h))

    def test_hero_photo_is_not_a_lockup(self):
        slide_w, slide_h = 20 * EMU_PER_INCH, int(11.25 * EMU_PER_INCH)
        slide_area = slide_w * slide_h
        left, top = int(10.1 * EMU_PER_INCH), 0
        width, height = int(9.9 * EMU_PER_INCH), slide_h
        self.assertFalse(is_brand_lockup(width * height, slide_area, left, top, width, height, slide_w, slide_h))


@unittest.skipUnless(TEMPLATE.exists(), "official template missing")
class TemplateGeometryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.prs = Presentation(str(TEMPLATE))
        cls.slide_w = int(cls.prs.slide_width)
        cls.slide_h = int(cls.prs.slide_height)

    def test_cover_frames_drop_logo_keep_hero(self):
        cover = self.prs.slides[0]
        frames = content_photo_frames(cover, self.slide_w, self.slide_h)
        self.assertEqual(len(frames), 1, "cover should expose one content photo well")
        self.assertGreater(frames[0]["width"], self.slide_w * 0.7, "cover photo is the full-bleed background")

    def test_prototype_insets_match_sample_slides(self):
        with tempfile.TemporaryDirectory() as tmp:
            tokens = inspect_template(TEMPLATE, Path(tmp), colorway="green")
        inset = tokens.get("content_inset") or {}
        self.assertGreaterEqual(inset.get("left", 0), int(0.8 * EMU_PER_INCH))
        self.assertLessEqual(inset.get("left", 0), int(1.4 * EMU_PER_INCH))
        cover = next(p for p in tokens["prototypes"] if p["role"] == "TITLE")
        box = content_safe_box(tokens, cover["index"])
        self.assertGreaterEqual(box["left"], inset.get("left", 0) - int(0.05 * EMU_PER_INCH))
        self.assertGreater(box["width"], int(8.0 * EMU_PER_INCH), "cover title sits on the card over the background photo")


def _write_source_with_photo(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(7.2), Inches(1.2))
    title.text_frame.paragraphs[0].text = "Field photo"
    body = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(6.5), Inches(4.5))
    body.text_frame.paragraphs[0].text = (
        "A longer body block beside a photograph so the slide stays a content layout. "
        "Staff must still be able to move or delete the picture in PowerPoint after convert."
    )
    slide.shapes.add_picture(io.BytesIO(_png_bytes((20, 90, 40), (400, 600))), Inches(8.4), Inches(0.3), Inches(4.4), Inches(6.9))
    prs.save(str(path))


def _write_title_and_section_no_photos(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    cover = prs.slides.add_slide(blank)
    title = cover.shapes.add_textbox(Inches(0.5), Inches(2.4), Inches(7.0), Inches(1.4))
    title.text_frame.paragraphs[0].text = "Claude Usage"
    subtitle = cover.shapes.add_textbox(Inches(0.5), Inches(4.0), Inches(7.0), Inches(0.8))
    subtitle.text_frame.paragraphs[0].text = "Executive Summary"
    section = prs.slides.add_slide(blank)
    heading = section.shapes.add_textbox(Inches(0.5), Inches(3.0), Inches(8.0), Inches(1.2))
    heading.text_frame.paragraphs[0].text = "Field Work"
    prs.save(str(path))


def _template_content_photo_hashes(slide_index: int) -> set[str]:
    prs = Presentation(str(TEMPLATE))
    frames = content_photo_frames(prs.slides[slide_index], int(prs.slide_width), int(prs.slide_height))
    return {f["hash"] for f in frames if f.get("hash")}


def _raster_hashes(slide, slide_w: int, slide_h: int) -> set[str]:
    """Picture shapes plus JPEG/PNG fills kept in the template wells."""
    found = set()
    for shape in _picture_shapes(slide):
        try:
            found.add(blob_hash(shape.image.blob))
        except Exception:
            continue
    for frame in content_photo_frames(slide, slide_w, slide_h):
        if frame.get("hash"):
            found.add(frame["hash"])
    return found


@unittest.skipUnless(TEMPLATE.exists() and SAMPLE.exists(), "template or sample fixture missing")
class ConvertedDeckTests(unittest.TestCase):
    def _assert_images_editable_and_clean(self, tokens, manifest, output: Path, source_hashes: set[str]):
        prs = Presentation(str(output))
        locked_source = []
        source_found = set()
        for slide in prs.slides:
            for shape in _picture_shapes(slide):
                digest = blob_hash(shape.image.blob)
                if digest in source_hashes:
                    source_found.add(digest)
                    if _pic_locks(shape):
                        locked_source.append(shape.name)
        self.assertFalse(locked_source, "source photographs must be unlocked (no picLocks)")
        self.assertEqual(source_found, source_hashes, "every source photograph must be re-inserted")

        # Text that isn't full-bleed stays on the template inset.
        inset_left = int((tokens.get("content_inset") or {}).get("left") or 0)
        if inset_left:
            for slide in prs.slides:
                for shape in slide.shapes:
                    if not getattr(shape, "has_text_frame", False):
                        continue
                    if not (shape.text_frame.text or "").strip():
                        continue
                    # Full-width titles are centered on the slide (left = 0) by design.
                    if int(shape.width or 0) >= int(prs.slide_width * 0.7):
                        continue
                    if int(shape.left or 0) < inset_left - int(0.15 * EMU_PER_INCH):
                        self.fail(f"{shape.name} left {shape.left} is inside the template margin {inset_left}")

    def test_sample_fixture_photos_unlocked_no_sample_art(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            result = _convert(TEMPLATE, SAMPLE, out)
            source_hashes = set()
            for slide in result["manifest"]["slides"]:
                for shape in slide["shapes"]:
                    if shape.get("kind") == "picture" and not shape.get("drop_logo"):
                        source_hashes.add(shape["image"]["hash"])
            self._assert_images_editable_and_clean(result["tokens"], result["manifest"], result["output"], source_hashes)
            marsh = next(e for e in result["plan"]["entries"] if e["source_index"] == 3)
            # Picture-heavy slide must not land on a no-photo statement layout.
            self.assertNotEqual(marsh.get("role"), "TITLE_ONLY")

    def test_right_photo_uses_photo_well_not_overlay(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "right-photo.pptx"
            _write_source_with_photo(src)
            result = _convert(TEMPLATE, src, Path(tmp) / "built")
            source_hashes = {
                shape["image"]["hash"]
                for slide in result["manifest"]["slides"]
                for shape in slide["shapes"]
                if shape.get("kind") == "picture" and not shape.get("drop_logo")
            }
            self._assert_images_editable_and_clean(result["tokens"], result["manifest"], result["output"], source_hashes)
            prs = Presentation(str(result["output"]))
            slide = prs.slides[0]
            pics = _picture_shapes(slide)
            source_pics = [p for p in pics if blob_hash(p.image.blob) in source_hashes]
            self.assertEqual(len(source_pics), 1)
            pic = source_pics[0]
            self.assertGreater(int(pic.left), int(prs.slide_width) * 0.4, "photo should occupy the right well, not cover left copy")
            self.assertFalse(_pic_locks(pic))
            # Newly added pictures must not sit on the title/body copy.
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                text = (shape.text_frame.text or "").strip()
                if not text or int(shape.left or 0) > int(prs.slide_width) * 0.45:
                    continue
                overlap_w = min(int(pic.left + pic.width), int(shape.left + shape.width)) - max(int(pic.left), int(shape.left))
                overlap_h = min(int(pic.top + pic.height), int(shape.top + shape.height)) - max(int(pic.top), int(shape.top))
                if overlap_w > 0 and overlap_h > 0:
                    frac = (overlap_w * overlap_h) / max(1, int(shape.width) * int(shape.height))
                    self.assertLess(frac, 0.25, f"photo covers text {text[:40]!r}")

    @unittest.skipUnless(KNIGHT.exists(), "Knight fixture not in this environment")
    def test_knight_fixture(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = _convert(TEMPLATE, KNIGHT, Path(tmp))
            source_hashes = {
                shape["image"]["hash"]
                for slide in result["manifest"]["slides"]
                for shape in slide["shapes"]
                if shape.get("kind") == "picture" and not shape.get("drop_logo")
            }
            self.assertEqual(len(source_hashes), 3)
            self._assert_images_editable_and_clean(result["tokens"], result["manifest"], result["output"], source_hashes)
            prs = Presentation(str(result["output"]))
            cover_pics = _picture_shapes(prs.slides[0])
            cover_hashes = {blob_hash(p.image.blob) for p in cover_pics}
            sample = set(result["tokens"].get("sample_photo_hashes") or [])
            self.assertTrue(cover_hashes & sample, "cover without a source photo keeps the title aerial")


@unittest.skipUnless(TEMPLATE.exists(), "official template missing")
class KeepTemplateStockPhotosTests(unittest.TestCase):
    def test_unused_wells_keep_title_and_section_stock(self):
        title_hashes = _template_content_photo_hashes(0)
        self.assertTrue(title_hashes, "template cover must have a background photograph")
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "no-photos.pptx"
            _write_title_and_section_no_photos(src)
            result = _convert(TEMPLATE, src, Path(tmp) / "built")
            prs = Presentation(str(result["output"]))
            self.assertGreaterEqual(len(prs.slides), 2)
            cover_hashes = _raster_hashes(prs.slides[0], int(prs.slide_width), int(prs.slide_height))
            self.assertTrue(
                title_hashes & cover_hashes,
                "converted slide 1 must keep the title-slide photograph when the source has no picture",
            )
            blob = "\n".join(
                (shape.text_frame.text or "").strip()
                for slide in prs.slides
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )
            for dummy in (
                "PRESENTATION TITLE",
                "SUBTITLE HERE",
                "NAME HERE",
                "SLIDE TITLE HERE",
                "Bullet 1",
                "Highlighted statement here",
            ):
                self.assertNotIn(dummy, blob)

    def test_source_picture_still_replaces_its_well(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "right-photo.pptx"
            _write_source_with_photo(src)
            result = _convert(TEMPLATE, src, Path(tmp) / "built")
            source_hashes = {
                shape["image"]["hash"]
                for slide in result["manifest"]["slides"]
                for shape in slide["shapes"]
                if shape.get("kind") == "picture" and not shape.get("drop_logo")
            }
            self.assertEqual(len(source_hashes), 1)
            prs = Presentation(str(result["output"]))
            found = {blob_hash(p.image.blob) for p in _picture_shapes(prs.slides[0])}
            self.assertTrue(source_hashes <= found, "mapped source photograph must replace the well")


class SanitizePictureTests(unittest.TestCase):
    def test_add_picture_locks_are_stripped(self):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        pic = slide.shapes.add_picture(io.BytesIO(_png_bytes((10, 20, 30))), Inches(1), Inches(1), Inches(2), Inches(2))
        self.assertTrue(_pic_locks(pic), "python-pptx defaults to noChangeAspect")
        sanitize_new_picture(pic)
        self.assertFalse(_pic_locks(pic))


if __name__ == "__main__":
    unittest.main(verbosity=2)
