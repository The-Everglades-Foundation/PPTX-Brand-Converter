#!/usr/bin/env python3
"""Regression: table (+ short caption) is the slide, not a clipped Comparison strip."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
sys.path.insert(0, str(ROOT))

from colorways import is_layout_chrome, is_short_caption  # noqa: E402
from helpers import data_content_box  # noqa: E402
from test_margins_images import TEMPLATE, _convert  # noqa: E402
from test_overlap_layers import _assert_no_dummy, _iter_text  # noqa: E402

EMU = 914400
# Source decks are 13.333″×7.5″; output follows the 20″×11.25″ template canvas.
SRC_W, SRC_H = 13.333, 7.5


def write_appendix_table_source(path: Path) -> None:
    """Copilot appendix analogue: title + caption + wide multi-row table."""
    prs = Presentation()
    prs.slide_width = Inches(SRC_W)
    prs.slide_height = Inches(SRC_H)
    blank = prs.slide_layouts[6]

    # Full-width table + caption (Appendix B page 1).
    s = prs.slides.add_slide(blank)
    title = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.2), Inches(0.55))
    title.text_frame.paragraphs[0].text = "Appendix B: All Copilot Users"
    cap = s.shapes.add_textbox(Inches(0.5), Inches(0.95), Inches(12.2), Inches(0.4))
    cap.text_frame.paragraphs[0].text = "Full member list, sorted by total prompts — users 1–22 of 22."
    header = ("Name", "Active Days", "Total Prompts", "Chat (work)", "Chat (web)")
    rows = [header]
    for i in range(1, 23):
        rows.append((f"User {i:02d}", str(i % 9 + 1), str(40 - i), str(i % 5), "0"))
    tbl = s.shapes.add_table(len(rows), 5, Inches(0.4), Inches(1.5), Inches(12.5), Inches(5.6)).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            tbl.cell(r, c).text = val

    # Two-column source: caption on the left, table on the right.
    s = prs.slides.add_slide(blank)
    title = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.2), Inches(0.55))
    title.text_frame.paragraphs[0].text = "Appendix A: All Claude Users"
    cap = s.shapes.add_textbox(Inches(0.4), Inches(1.5), Inches(3.6), Inches(4.8))
    cap.text_frame.paragraphs[0].text = "Full member list, sorted by chats — 12 of 12 licensed seats."
    header = ("Name", "Chats", "Projects", "Tools")
    rows = [header]
    for i in range(1, 13):
        rows.append((f"Member {i:02d}", str(20 + i), str(i % 4), str(i % 3)))
    tbl = s.shapes.add_table(len(rows), 4, Inches(4.3), Inches(1.4), Inches(8.5), Inches(5.5)).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            tbl.cell(r, c).text = val

    prs.save(str(path))


def _shape_box(shape) -> tuple[int, int, int, int]:
    return (
        int(shape.left or 0),
        int(shape.top or 0),
        int(shape.width or 0),
        int(shape.height or 0),
    )


def _run_rgb(shape) -> str | None:
    try:
        run = shape.text_frame.paragraphs[0].runs[0]
        return str(run.font.color.rgb)
    except Exception:
        return None


class HelperTests(unittest.TestCase):
    def test_short_caption_detects_appendix_line(self):
        self.assertTrue(
            is_short_caption("Full member list, sorted by total prompts — users 24–46 of 46.")
        )
        self.assertFalse(is_short_caption("COMPARISON"))
        long_copy = " ".join(["Teachers ask in plain English"] * 12)
        self.assertFalse(is_short_caption(long_copy))

    def test_data_content_box_reaches_the_floor(self):
        tokens = {"slide_width": int(20 * EMU), "slide_height": int(11.25 * EMU)}
        box = data_content_box(tokens, None)
        bottom = int(box["top"]) + int(box["height"])
        self.assertGreater(int(box["height"]), int(5.5 * EMU))
        self.assertLessEqual(bottom, int(tokens["slide_height"]))
        self.assertGreaterEqual(bottom, int(10.5 * EMU))


@unittest.skipUnless(TEMPLATE.exists(), "official template missing")
class TableOnSlideTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.src = Path(cls.tmp.name) / "appendix_table.pptx"
        write_appendix_table_source(cls.src)
        cls.result = _convert(TEMPLATE, cls.src, Path(cls.tmp.name) / "built")
        cls.prs = Presentation(str(cls.result["output"]))
        cls.plan = cls.result["plan"]
        cls.slide_w = int(cls.prs.slide_width)
        cls.slide_h = int(cls.prs.slide_height)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _table_slides(self, needle: str) -> list:
        found = []
        for slide in self.prs.slides:
            blob = "\n".join(_iter_text(slide))
            if needle not in blob:
                continue
            if any(s.shape_type == MSO_SHAPE_TYPE.TABLE for s in slide.shapes):
                found.append(slide)
        self.assertTrue(found, f"no table slide contains {needle!r}")
        return found

    def _assert_table_on_canvas(self, shape, label: str) -> None:
        left, top, width, height = _shape_box(shape)
        self.assertGreaterEqual(left, 0, label)
        self.assertGreaterEqual(top, 0, label)
        self.assertLessEqual(left + width, self.slide_w + int(0.02 * EMU), f"{label} overflows right")
        self.assertLessEqual(top + height, self.slide_h + int(0.02 * EMU), f"{label} overflows bottom")
        self.assertGreater(
            height,
            int(3.0 * EMU),
            f"{label} is a thin bottom strip ({height / EMU:.2f}″), not the content well",
        )
        self.assertLess(
            top,
            int(self.slide_h * 0.45),
            f"{label} starts too low ({top / EMU:.2f}″) — parked in a leftover banner",
        )
        self.assertGreaterEqual(
            height / max(1, self.slide_h),
            0.32,
            f"{label} must occupy the main well under the title",
        )

    def test_appendix_b_table_fills_the_well(self):
        slides = self._table_slides("Appendix B: All Copilot Users")
        for i, slide in enumerate(slides):
            _assert_no_dummy(slide, f"appendix B {i}")
            blob = "\n".join(_iter_text(slide))
            self.assertNotIn("COLUMN ONE", blob)
            self.assertNotIn("Row label", blob)
            tables = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TABLE]
            self.assertTrue(tables)
            for table in tables:
                self._assert_table_on_canvas(table, f"Appendix B table {i}")
            captions = [
                s
                for s in slide.shapes
                if getattr(s, "has_text_frame", False)
                and "Full member list" in (s.text_frame.text or "")
            ]
            self.assertTrue(captions, "caption must stay with the table")
            for cap in captions:
                _l, _t, _w, h = _shape_box(cap)
                self.assertLess(h, int(1.6 * EMU), "caption must not swallow the cream well")
                color = _run_rgb(cap)
                self.assertNotEqual(color, "FFFFFF", "caption must not be white-on-cream")
                self.assertIn("Full member list", cap.text_frame.text)

    def test_two_column_source_still_places_table_as_the_slide(self):
        slides = self._table_slides("Appendix A: All Claude Users")
        for i, slide in enumerate(slides):
            _assert_no_dummy(slide, f"appendix A {i}")
            tables = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TABLE]
            self.assertTrue(tables)
            for table in tables:
                self._assert_table_on_canvas(table, f"Appendix A table {i}")
            blob = "\n".join(_iter_text(slide))
            self.assertIn("Full member list", blob)
            # Caption + table belong together — not a title-only THE POINT then a strip.
            self.assertFalse(
                "THE POINT" in blob and "COMPARISON" not in blob and "Full member list" in blob and not tables
            )

    def test_plan_keeps_caption_on_the_table_slide(self):
        for entry in self.plan["entries"]:
            kinds = {p.get("kind") for p in entry.get("placements") or []}
            if "table" not in kinds:
                continue
            name = entry.get("template_layout_name") or ""
            self.assertTrue(
                "Comparison" in name or "comparison" in name.lower() or "Data" in name,
                f"table slide should use Comparison, got {name!r}",
            )


if __name__ == "__main__":
    unittest.main(verbosity=2)
